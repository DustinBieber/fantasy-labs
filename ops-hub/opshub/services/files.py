"""Read spreadsheets, PDFs and text files that arrive by email, Drive or upload."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime

TABLE_EXTS = (".csv", ".tsv", ".xlsx", ".xlsm")
TEXT_EXTS = (".txt", ".md", ".json", ".csv", ".tsv")


def _decode(content: str | bytes) -> str:
    if isinstance(content, str):
        return content
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(filename: str, content: str | bytes) -> list[list[str]]:
    """Rows of a CSV/TSV or the first non-empty sheet of an Excel workbook, as strings."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        data = content if isinstance(content, bytes) else content.encode()
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            if rows:
                return rows
        return []
    text = _decode(content)
    delim = "\t" if name.endswith(".tsv") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def table_to_csv(rows: list[list[str]]) -> str:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue()


def pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def to_text(filename: str, content: str | bytes) -> str:
    """Best-effort plain text of any supported file (for AI reading or display)."""
    name = filename.lower()
    if name.endswith(".pdf") and isinstance(content, bytes):
        return pdf_text(content)
    if name.endswith((".xlsx", ".xlsm")):
        return table_to_csv(read_table(filename, content))
    return _decode(content)
