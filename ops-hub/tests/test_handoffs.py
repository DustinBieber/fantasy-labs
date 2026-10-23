"""RQ report imports, Muse handoff, PDF/Excel invoices, chat commands, scheduler, migrations."""
import hashlib
import hmac
import io
import json
from datetime import timedelta

import httpx
from sqlalchemy import select

from opshub.connectors.base import Connector, NDocument
from opshub.connectors.iqmetrix_rq import IQMetrixRQConnector, parse_report
from opshub.models import Alert, Forecast, HandoffFile, InboxItem, Invoice, Order, Shipment, StockLevel
from opshub.services import chatops, documents, jobs, muse, rules
from opshub.timeutil import utcnow


def _xlsx(rows) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- RQ ------------------------------------------------------------------------------------

def test_rq_inventory_report_with_title_rows_and_totals():
    rows = [["RQ - Inventory Listing"], ["Printed 10/23/2026"], ["Store Name", "Product SKU", "Description", "Qty On Hand",
                                                                "Committed"],
            ["Eastgate", "PWR-20K", "Power bank", "4", "1"], ["Eastgate", "PWR-20K", "Power bank", "2", ""],
            ["", "Total", "", "6", ""]]
    kind, b = parse_report(rows, default_store="RQ", tz="Pacific/Honolulu")
    assert kind == "inventory" and len(b.stock) == 1
    assert (b.stock[0].location, b.stock[0].on_hand, b.stock[0].available) == ("Eastgate", 6, 5)


def test_rq_sales_report_from_excel():
    data = _xlsx([["Invoice #", "Date", "Store", "SKU", "Qty Sold", "Extended Price"],
                  ["E-1", "10/22/2026 3:45 PM", "Eastgate", "PWR-20K", 1, 39.99],
                  ["E-2", "10/21/2026", "Westfield", "CASE-UNI", 2, 39.98]])
    batch, report = IQMetrixRQConnector(env={}).import_files([("sales.xlsx", data)], "Pacific/Honolulu")
    assert report[0]["kind"] == "sales" and len(batch.sales) == 2
    assert batch.sales[0].sold_at.hour == 1  # 3:45 PM Honolulu is 01:45 UTC next day


def test_rq_unknown_file_is_reported_not_imported():
    batch, report = IQMetrixRQConnector(env={}).import_files([("notes.csv", "hello,world\n1,2\n")], "UTC")
    assert report[0]["kind"] == "unknown" and not batch.stock


def test_rq_upload_endpoint(synced, settings):
    from fastapi.testclient import TestClient

    from opshub.api import create_app

    with TestClient(create_app()) as c:
        csv_text = "Store,SKU,Description,On Hand\n101,PWR-20K,Power bank,40\n"
        r = c.post("/api/import/rq?filename=inv.csv", content=csv_text.encode()).json()
        assert r["imported"] and r["files"][0]["kind"] == "inventory"
    level = synced.scalar(select(StockLevel).join(StockLevel.location).where(StockLevel.product.has(sku="PWR-20K"))
                          .where(StockLevel.location.has(code="RQ-101")))
    synced.refresh(level)
    assert level.on_hand == 40  # store "101" maps to RQ-101 through the Locations sheet


# --- Muse ------------------------------------------------------------------------------------

def test_muse_import_creates_items_deadlines_and_forecasts(synced):
    titles = {i.title for i in synced.scalars(select(InboxItem))}
    assert "Holiday Gift Box selling slower than plan" in titles
    o = synced.scalar(select(Order).where(Order.number == "#2054"))
    assert o.required_by_source == "muse" and "note from muse" in o.note_flags
    assert {f.product.sku for f in synced.scalars(select(Forecast))} == {"PWR-20K", "CBC-32"}
    kinds = {a.kind for a in synced.scalars(select(Alert).where(Alert.status == "open"))}
    assert {"muse_alert", "muse_task", "muse_note"} <= kinds


def test_muse_new_file_supersedes_old_alerts_and_csv_works(synced, settings):
    csv_text = "type,severity,title,sku\nalert,urgent,Reorder power banks now,PWR-20K\ntask,,Call the 3PL,\n"
    muse.import_document(synced, NDocument(external_id="m2", title="m2.csv", text="",
                                           attachments=[("muse-update.csv", csv_text)], hint="muse"), settings)
    alerts = {i.title: i.status for i in synced.scalars(select(InboxItem).where(InboxItem.kind == "alert"))}
    assert alerts == {"Holiday Gift Box selling slower than plan": "superseded", "Reorder power banks now": "open"}
    rules.run(synced, settings)
    a = synced.scalar(select(Alert).where(Alert.title == "Reorder power banks now"))
    assert a.severity == "urgent" and a.subject_type == "product"


def test_muse_free_text_becomes_a_note(synced, settings):
    muse.import_document(synced, NDocument(external_id="m3", title="thoughts.md", text="",
                                           attachments=[("thoughts.md", "Push matcha this week.\nMore detail.")],
                                           hint="muse"), settings, env={})
    note = synced.scalar(select(InboxItem).where(InboxItem.file_name == "thoughts.md"))
    assert note.kind == "note" and note.title == "Muse note: Push matcha this week."


def test_muse_export_writes_brief_and_data(synced, settings, tmp_path):
    result = muse.export_daily(synced, settings, env={"MUSE_LOCAL_DIR": str(tmp_path)})
    data = json.loads((tmp_path / "to-muse" / result["files"][0]).read_text())
    assert data["kind"] == "opshub_to_muse" and data["orders_watch"] and data["inventory"]["sell_through"]
    brief = (tmp_path / "to-muse" / result["files"][1]).read_text()
    assert brief.startswith("# Ops Hub brief") and "## Freight in transit" in brief
    assert synced.scalar(select(HandoffFile).where(HandoffFile.direction == "out")) is not None


# --- invoices from Excel / PDF -------------------------------------------------------------------

def test_excel_invoice_and_unread_pdf(session, settings):
    rows = [["Vendor", "Riverside 3PL"], ["Invoice Number", "X-1"], ["Billing Period", "2026-09-01", "2026-09-30"],
            [], ["Charge Code", "Description", "Quantity", "Rate", "Amount"], ["PICK-ORD", "Picks", 2, 2.75, 5.5]]
    kind = documents.process(session, NDocument(external_id="e1", title="inv", text="",
                                                attachments=[("inv.xlsx", _xlsx(rows))]), "gmail", settings, env={})
    assert kind == "invoice" and session.scalar(select(Invoice)).number == "X-1"

    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    # A PDF with no "invoice" text is ignored; that's what a blank page gives us.
    assert documents.process(session, NDocument(external_id="p1", title="scan", text="",
                                                attachments=[("scan.pdf", buf.getvalue())]), "gmail", settings,
                             env={}) == "ignored"


# --- chat commands ---------------------------------------------------------------------------------

def test_chat_roles():
    env = {"TELEGRAM_CHAT_ID": "-100", "TELEGRAM_FREIGHT_CHAT_IDS": "55", "WHATSAPP_TO": "+15550001",
           "WHATSAPP_FREIGHT_NUMBERS": "15550099"}
    assert chatops.role_for("telegram", "-100", env) == "staff"
    assert chatops.role_for("telegram", "55", env) == "freight"
    assert chatops.role_for("telegram", "12", env) is None
    assert chatops.role_for("whatsapp", "15550001", env) == "staff"  # + is optional
    assert chatops.role_for("whatsapp", "15550099", env) == "freight"


def test_chat_commands_and_carrier_updates(synced, settings):
    assert "Order #2041" in chatops.handle(synced, settings, "order 2041", "staff", "telegram")
    assert "HQ-WH" in chatops.handle(synced, settings, "/stock pwr-20k", "staff", "telegram")
    assert "Central Plains" in chatops.handle(synced, settings, "ship 4471983", "staff", "telegram")
    # Freight contacts can't query, but can update by quoting a reference.
    assert "include the PRO" in chatops.handle(synced, settings, "stock PWR-20K", "freight", "whatsapp")
    reply = chatops.handle(synced, settings, "RW-30017 handed off to Prairie Line, new PRO PL-77. ETA 10/30/2026",
                           "freight", "whatsapp")
    assert reply.startswith("Logged on TR-HQ-0921") and "Prairie Line" in reply
    shp = synced.scalar(select(Shipment).where(Shipment.reference == "TR-HQ-0921"))
    assert shp.legs[-1].tracking_ref == "PL-77" and shp.events[-1].source == "whatsapp"


def test_telegram_poll_answers_staff_only(synced, settings, monkeypatch):
    sent = []

    def handler(req: httpx.Request):
        if req.url.path.endswith("/getUpdates"):
            return httpx.Response(200, json={"result": [
                {"update_id": 10, "message": {"chat": {"id": -100}, "text": "alerts"}},
                {"update_id": 11, "message": {"chat": {"id": 999}, "text": "alerts"}}]})
        sent.append(json.loads(req.content))
        return httpx.Response(200, json={"ok": True})

    monkeypatch.setattr(Connector, "transport", httpx.MockTransport(handler))
    env = {"TELEGRAM_BOT_TOKEN": "t", "TELEGRAM_CHAT_ID": "-100"}
    assert chatops.poll_telegram(synced, settings, env) == {"updates": 2, "replied": 1}
    assert sent[0]["chat_id"] == "-100" and "open" in sent[0]["text"]
    from opshub.services import kv

    assert kv.get("telegram_offset", session=synced) == "12"


def test_whatsapp_webhook_signature(synced, settings, monkeypatch):
    from fastapi.testclient import TestClient

    from opshub.api import create_app

    monkeypatch.setenv("WHATSAPP_APP_SECRET", "shh")
    monkeypatch.setenv("WHATSAPP_VERIFY_TOKEN", "vt")
    monkeypatch.setenv("WHATSAPP_FREIGHT_NUMBERS", "15550099")
    body = json.dumps({"entry": [{"changes": [{"value": {"messages": [
        {"from": "15550099", "type": "text", "text": {"body": "PRO CPF-220931 delivered"}}]}}]}]}).encode()
    sig = "sha256=" + hmac.new(b"shh", body, hashlib.sha256).hexdigest()
    with TestClient(create_app()) as c:
        assert c.get("/webhooks/whatsapp", params={"hub.mode": "subscribe", "hub.verify_token": "vt",
                                                   "hub.challenge": "123"}).text == "123"
        assert c.post("/webhooks/whatsapp", content=body, headers={"X-Hub-Signature-256": "sha256=bad"}).status_code == 401
        assert c.post("/webhooks/whatsapp", content=body, headers={"X-Hub-Signature-256": sig}).json()["ok"]
    shp = synced.scalar(select(Shipment).where(Shipment.reference == "OUT-WH-88213"))
    synced.refresh(shp)
    assert shp.status == "delivered"


# --- scheduler, migrations -------------------------------------------------------------------------

def test_scheduler_runs_sync_and_daily_jobs_once(session, settings, monkeypatch, tmp_path):
    from opshub.config import set_settings

    monkeypatch.chdir(tmp_path)  # Muse demo export writes ./handoff
    settings.sync_interval_minutes, settings.digest_time, settings.muse_export_time = 15, "00:00", "00:00"
    set_settings(settings)
    sched = jobs.Scheduler(env={})
    now = utcnow()
    first = sched.tick(now)
    assert first[:1] == ["sync"] and {"muse_export", "digest"} <= set(first)
    assert sched.tick(now + timedelta(minutes=5)) == []  # nothing due yet
    assert sched.tick(now + timedelta(minutes=16)) == ["sync", "calendar", "sheets_report"]
    assert (tmp_path / "handoff" / "to-muse").is_dir()


def test_old_database_gets_new_columns(tmp_path):
    import sqlite3

    from opshub import db

    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE orders (id INTEGER PRIMARY KEY, source VARCHAR(40), external_id VARCHAR(120), "
                "number VARCHAR(120), status VARCHAR(32), order_type VARCHAR(32), synced_at DATETIME)")
    con.commit()
    con.close()
    db.init_engine(f"sqlite:///{path}")
    db.create_all()
    cols = {r[1] for r in sqlite3.connect(path).execute("pragma table_info(orders)")}
    assert {"manual_deadline", "external_notes", "shipped_at", "required_by"} <= cols
